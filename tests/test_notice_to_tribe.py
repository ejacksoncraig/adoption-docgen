"""The Notice to Tribe: generated only when ICWA applies, addressed to whichever
tribe was named, for either matter and either child count."""

from __future__ import annotations

import docx
import pytest
from docx.shared import Inches

from app import engine
from tests import fixtures
from tests.test_render import render

NOTICE = "notice_to_tribe.docx"

KNOWN_TRIBES = {
    "Cherokee Nation": ("ICW Adoptions", "P.O. Box 948", "Tahlequah, OK 74465"),
    "Chickasaw Nation": ("Adoption Unit", "1401 Hoppe Blvd", "Ada, OK 74820"),
    "Choctaw Nation of Oklahoma": ("ICW Adoptions", "P.O. Box 1210", "Durant, OK 74702-1210"),
    "Muscogee (Creek) Nation": ("Children and Family Services", "P.O. Box 580", "Okmulgee, OK 74447"),
    "Citizen Potawatomi Nation": ("ICW / Adoptions", "1601 S. Gordon Cooper Drive", "Shawnee, OK 74801"),
}


# --------------------------------------------------------------------------
# generated only when ICWA applies
# --------------------------------------------------------------------------


def test_notice_is_generated_when_icwa_applies(registry, tmp_path, today):
    result = engine.generate(registry, "dhs", "dhs_1p_1c", fixtures.BASE, today=today, output_root=tmp_path)
    assert any(f.template.endswith(NOTICE) for f in result.files)


def test_notice_is_skipped_when_icwa_does_not_apply(registry, tmp_path, today):
    result = engine.generate(registry, "dhs", "dhs_1p_1c", fixtures.ALTERNATE, today=today, output_root=tmp_path)
    assert not any(f.template.endswith(NOTICE) for f in result.files)


def test_notice_is_generated_for_a_stepparent_filing_when_icwa_applies(registry, tmp_path, today):
    values = fixtures.stepparent(icwa_applies=True, tribe="Cherokee Nation")
    result = engine.generate(registry, "stepparent", "step_1c", values, today=today, output_root=tmp_path)
    assert any(f.template.endswith(NOTICE) for f in result.files)


def test_notice_is_skipped_for_a_stepparent_filing_when_icwa_does_not_apply(registry, tmp_path, today):
    result = engine.generate(registry, "stepparent", "step_1c", fixtures.STEPPARENT, today=today, output_root=tmp_path)
    assert not any(f.template.endswith(NOTICE) for f in result.files)


# --------------------------------------------------------------------------
# no leftover placeholders, whichever branch
# --------------------------------------------------------------------------


@pytest.mark.parametrize("token", ["XXX", "{{", "{%"])
def test_no_placeholder_survives_into_the_notice(registry, tmp_path, today, token):
    text = render(registry, fixtures.BASE, tmp_path, today, template=NOTICE)
    assert token not in text


# --------------------------------------------------------------------------
# the correct tribe's office, addressed automatically
# --------------------------------------------------------------------------


@pytest.mark.parametrize("tribe,expected", KNOWN_TRIBES.items())
def test_each_known_tribe_gets_its_own_office_address(registry, tmp_path, today, tribe, expected):
    text = render(registry, fixtures.values(tribe=tribe), tmp_path, today, template=NOTICE)
    assert tribe in text
    for line in expected:
        assert line in text
    other_addresses = [addr for name, addr in KNOWN_TRIBES.items() if name != tribe]
    for _, _, city in other_addresses:
        assert city not in text


def test_a_tribe_outside_the_known_list_still_generates_a_notice(registry, tmp_path, today):
    """The office told us to expect more tribes than these five. A name that
    doesn't match one of them still produces a notice — it just says so, rather
    than printing an address that belongs to a different tribe."""
    text = render(registry, fixtures.values(tribe="Osage Nation"), tmp_path, today, template=NOTICE)
    assert "Osage Nation" in text
    assert "Tribe mailing address not on file" in text
    for _, _, city in KNOWN_TRIBES.values():
        assert city not in text


# --------------------------------------------------------------------------
# one child or several, in the caption
# --------------------------------------------------------------------------


def test_single_child_caption_reads_in_the_singular(registry, tmp_path, today):
    text = render(registry, fixtures.BASE, tmp_path, today, template=NOTICE)
    assert "A MINOR CHILD" in text
    assert "MINOR CHILDREN" not in text
    assert "is a member of the Sample Nation" in text
    assert "And" not in text.split("NOTICE TO TRIBE")[0]


def test_two_children_share_one_notice_in_the_plural(registry, tmp_path, today):
    values = {
        **fixtures.BASE,
        "child2_name": "SAMPLE CHILD TWO",
        "child2_dob": "2016-02-15",
        "child2_race": "Caucasian",
        "child2_gender": "female",
        "child2_birth_place": "Sample Regional Hospital",
        "child2_birth_city": "Sampletown",
        "child2_birth_county": "Sample",
        "child2_birth_state": "Oklahoma",
        "name_change_child2": False,
    }
    text = render(registry, values, tmp_path, today, template=NOTICE, variant="dhs_1p_2c")
    assert "MINOR CHILDREN" in text
    assert "A MINOR CHILD" not in text
    assert "SAMPLE CHILD" in text and "SAMPLE CHILD TWO" in text
    assert "are members of the Sample Nation" in text
    assert 'are each an "Indian Child"' in text


# --------------------------------------------------------------------------
# the hearing paragraph
# --------------------------------------------------------------------------


HEARING_ANSWERS = {
    "petition_hearing_date": "2026-06-18",
    "courthouse_city": "Wagoner",
    "judge_name": "Douglas Kirkley",
}


def test_the_hearing_paragraph_prints_blanks_until_it_is_scheduled(registry, tmp_path, today):
    """All four answers are optional: a notice is often drafted before the
    hearing is set. Left blank, ¶1 reads exactly as the office's own form does."""
    text = render(registry, fixtures.BASE, tmp_path, today, template=NOTICE)
    assert "heard on the __________ day of ____________, 20____," in text
    assert "before the Honorable Judge ________________" in text
    assert "in the District Courthouse at ______________, Wagoner County" in text


def test_the_hearing_paragraph_is_filled_in_when_it_is_known(registry, tmp_path, today):
    text = render(registry, fixtures.values(**HEARING_ANSWERS), tmp_path, today, template=NOTICE)
    assert "heard on the 18th day of June, 2026, at" in text
    assert "before the Honorable Judge Douglas Kirkley" in text
    assert "in the District Courthouse at Wagoner, Wagoner County, State of Oklahoma." in text


def test_the_courthouse_city_is_asked_for_separately_from_the_county(registry, tmp_path, today):
    """The courthouse is in the county seat, which is not the county's name —
    Wagoner County's is Wagoner, but Kay County's is Newkirk."""
    values = fixtures.values(county="Kay", courthouse_city="Newkirk")
    text = render(registry, values, tmp_path, today, template=NOTICE)
    assert "in the District Courthouse at Newkirk, Kay County" in text


# --------------------------------------------------------------------------
# the birth parents, ¶3
# --------------------------------------------------------------------------


def test_the_birth_parents_are_named_in_a_dhs_notice(registry, tmp_path, today):
    text = render(registry, fixtures.BASE, tmp_path, today, template=NOTICE)
    assert (
        "the names of the birth parents of the minor child(ren) are "
        "SAMPLE BIRTH MOTHER, and SAMPLE BIRTH FATHER."
    ) in text


def test_the_birth_parents_of_a_stepparent_filing_are_the_petitioner_and_the_other_parent(
    registry, tmp_path, today
):
    """A step-parent matter has no bio_parents group: petitioner 1 *is* a birth
    parent and the `stepparent` group holds the other one. Both are optional to
    this document, so each side guards its own."""
    values = fixtures.stepparent(icwa_applies=True, tribe="Cherokee Nation")
    text = render(registry, values, tmp_path, today, template=NOTICE,
                  matter="stepparent", variant="step_1c")
    assert (
        "the names of the birth parents of the minor child(ren) are "
        "SAMPLE PARENT, and SAMPLE OTHER PARENT."
    ) in text


# --------------------------------------------------------------------------
# the certificate of service
# --------------------------------------------------------------------------


def test_the_mailing_date_prints_blanks_until_the_notice_goes_out(registry, tmp_path, today):
    text = render(registry, fixtures.BASE, tmp_path, today, template=NOTICE)
    assert "return receipt requested, on the ____ day of __________, 20____," in text


def test_the_mailing_date_is_filled_in_when_it_is_known(registry, tmp_path, today):
    values = fixtures.values(notice_mailing_date="2026-05-04")
    text = render(registry, values, tmp_path, today, template=NOTICE)
    assert "return receipt requested, on the 4th day of May, 2026, and receipt" in text


# --------------------------------------------------------------------------
# the signature block sits on the right
# --------------------------------------------------------------------------


#: Where the office put the signature block: the "By____" line at 3 inches, the
#: name and address beneath it a quarter-inch further in, so the text sits under
#: the signature line rather than under the word "By".
SIGNATURE_LINE = Inches(3)
SIGNATURE_BLOCK = Inches(3.25)


def _notice_document(registry, tmp_path, today, values=None, variant="dhs_1p_1c"):
    result = engine.generate(
        registry, "dhs", variant, values or fixtures.BASE,
        today=today, output_root=tmp_path,
    )
    notice = next(f for f in result.files if f.template.endswith(NOTICE))
    return docx.Document(str(notice.path))


def test_both_signature_blocks_sit_on_the_right(registry, tmp_path, today):
    document = _notice_document(registry, tmp_path, today)

    by_lines, block_lines = [], []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("By____"):
            by_lines.append(paragraph)
        elif text.startswith(("Attorney for Petitioner", "P. O. Box", "Wagoner, OK")) or "OBA #" in text:
            block_lines.append(paragraph)

    # Two blocks: one under the notice, one under the certificate of service.
    assert len(by_lines) == 2
    assert len(block_lines) == 8
    for paragraph in by_lines:
        assert paragraph.paragraph_format.left_indent == SIGNATURE_LINE, paragraph.text
    for paragraph in block_lines:
        assert paragraph.paragraph_format.left_indent == SIGNATURE_BLOCK, paragraph.text


# --------------------------------------------------------------------------
# the caption the office formatted by hand
# --------------------------------------------------------------------------


def caption_table(document):
    """The caption's 1x3 layout table: party text, parentheses, case number."""
    from docx.oxml.ns import qn
    from docx.table import Table

    for child in document.element.body.iterchildren():
        if child.tag != qn("w:tbl"):
            continue
        grid = child.find(qn("w:tblGrid"))
        if grid is not None and len(grid) == 3:
            return Table(child, document)
    raise AssertionError("no three-column caption table in the document")


def caption_lines(document):
    """Each caption row as (party text, parenthesis, case number)."""
    cells = caption_table(document).rows[0].cells
    depth = max(len(c.paragraphs) for c in cells)
    rows = []
    for i in range(depth):
        rows.append(tuple(
            c.paragraphs[i].text if i < len(c.paragraphs) else "" for c in cells
        ))
    return rows


def test_the_caption_keeps_every_parenthesis_in_its_own_column(registry, tmp_path, today):
    """The parentheses used to be tabbed out, which only lined up while the text
    beside them stayed short — a long child's name pushed its own ")" to the next
    stop and broke the rule. A column cannot drift, whatever the name."""
    document = _notice_document(registry, tmp_path, today, fixtures.values(
        child1_name="MAXIMILIAN BARTHOLOMEW CUNNINGHAM-WHITTINGTON"))
    parentheses = [paren for _text, paren, _number in caption_lines(document)]
    assert parentheses == [")"] * 5


def test_the_caption_names_the_child_and_the_case_in_their_own_columns(
    registry, tmp_path, today
):
    rows = caption_lines(_notice_document(registry, tmp_path, today))
    assert rows[0][0] == "IN THE MATTER OF THE ADOPTION OF:"
    assert rows[2][0] == "SAMPLE CHILD\t\tDOB:  3/2/2015"
    assert rows[2][2] == "Case No.  FA-2026-_____"
    # The case number belongs in the third column and only there.
    assert [n for _t, _p, n in rows if n.strip()] == ["Case No.  FA-2026-_____"]


def test_only_the_minor_child_line_is_indented(registry, tmp_path, today):
    """Half an inch on the "A MINOR CHILD" line, nothing on the rest."""
    cell = caption_table(_notice_document(registry, tmp_path, today)).rows[0].cells[0]
    indented = {p.text.strip(): p.paragraph_format.first_line_indent for p in cell.paragraphs}
    assert indented["A MINOR CHILD"] == Inches(0.5)
    assert all(v is None for k, v in indented.items() if k != "A MINOR CHILD")


def test_a_second_child_is_added_to_the_caption_without_moving_the_column(
    registry, tmp_path, today
):
    values = {
        **fixtures.BASE,
        "child2_name": "SAMPLE CHILD TWO",
        "child2_dob": "2016-02-15",
        "child2_race": "Caucasian",
        "child2_gender": "female",
        "child2_birth_place": "Sample Regional Hospital",
        "child2_birth_city": "Sampletown",
        "child2_birth_county": "Sample",
        "child2_birth_state": "Oklahoma",
        "name_change_child2": False,
    }
    rows = caption_lines(
        _notice_document(registry, tmp_path, today, values, variant="dhs_1p_2c"))

    # Two more lines than the singular caption: the "And" and the second child.
    assert [paren for _t, paren, _n in rows] == [")"] * 7
    assert rows[3][0] == "And"
    assert "SAMPLE CHILD TWO" in rows[4][0]
    assert [n for _t, _p, n in rows if n.strip()] == ["Case No.  FA-2026-_____"]


def test_the_caption_table_prints_no_rules(registry, tmp_path, today):
    from docx.oxml.ns import qn

    table = caption_table(_notice_document(registry, tmp_path, today))
    borders = table._tbl.find(qn("w:tblPr")).find(qn("w:tblBorders"))
    assert borders is not None, "the caption table declares no borders at all"
    assert {edge.get(qn("w:val")) for edge in borders} == {"none"}


# --------------------------------------------------------------------------
# the three address blocks agree
# --------------------------------------------------------------------------


def test_every_address_block_carries_the_same_four_lines(registry, tmp_path, today):
    """The notice addresses the tribe three times — once at the top, twice in the
    certificate of service, the latter two inside layout tables. All three carry
    the department line; one lost it when the table was first typed up."""
    text = render(registry, fixtures.values(tribe="Cherokee Nation"), tmp_path, today,
                  template=NOTICE)
    for line in ("ICW Adoptions", "P.O. Box 948", "Tahlequah, OK 74465"):
        assert text.count(line) == 3, f"{line!r} appears {text.count(line)} times, expected 3"
    # The tribe's name also appears in the paragraph that states membership.
    assert text.count("Cherokee Nation") == 4
