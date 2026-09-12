"""The attorney's signature image: what is accepted, where it lands, who it signs for.

The last of those is the one that matters. This program puts a signature on the
page on someone's behalf, so the test that earns its keep is not "the image was
embedded" — it is that the image was embedded *above the attorney's own block*
and nowhere near the petitioners', the notary's or the judge's line.
"""

from __future__ import annotations

import math
import re
import struct
import zipfile
import zlib
from pathlib import Path

import pytest

from app import engine, signature
from app.registry import SIGNATURE_SETTING, Registry
from tests import fixtures

PARAGRAPH = re.compile(r'<w:p\b(?:(?!</w:p>).)*?</w:p>', re.S)


# --------------------------------------------------------------------------
# a stand-in signature
# --------------------------------------------------------------------------


def png(width: int = 900, height: int = 260) -> bytes:
    """A PNG of about a signature's proportions, written without Pillow.

    A real scan is not committed here for the obvious reason, and a solid
    rectangle would not exercise the aspect-ratio scaling that keeps a long thin
    signature from running off the end of its line.
    """
    rows = []
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            ink = abs(y - height * (0.55 + 0.3 * math.sin(x / width * 11)))
            row += bytes((10, 12, 60, int(255 * max(0.0, 1.0 - ink / 5.0))))
        rows.append(bytes(row))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 6, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"".join(rows), 1))
            + chunk(b"IEND", b""))


@pytest.fixture
def config(tmp_path) -> Path:
    folder = tmp_path / "config"
    folder.mkdir()
    return folder


@pytest.fixture
def signed_registry(registry, tmp_path):
    """The real registry, told there is a signature on file.

    The image is written under tmp_path rather than into the repository's own
    config/, so a test run never leaves a signature lying about on disk.
    """
    image = tmp_path / "signature.png"
    image.write_bytes(png())
    registry.settings[SIGNATURE_SETTING] = str(image)
    try:
        yield registry
    finally:
        registry.settings.pop(SIGNATURE_SETTING, None)


# --------------------------------------------------------------------------
# what is accepted
# --------------------------------------------------------------------------


def test_a_png_is_stored_under_its_real_extension(config):
    path = signature.save(png(), config)
    assert path.name == "signature.png"
    assert signature.stored_path(config) == path


def test_a_file_that_is_not_an_image_is_refused(config):
    with pytest.raises(signature.SignatureError) as caught:
        signature.save(b"%PDF-1.4\nnot a picture", config)
    assert "not an image" in caught.value.problems[0]
    assert signature.stored_path(config) is None


def test_an_empty_file_is_refused(config):
    with pytest.raises(signature.SignatureError):
        signature.save(b"", config)


def test_a_photograph_of_a_whole_page_is_refused_by_size(config):
    with pytest.raises(signature.SignatureError) as caught:
        signature.save(b"\x89PNG\r\n\x1a\n" + b"\0" * signature.MAX_BYTES, config)
    assert "larger than a signature scan" in caught.value.problems[0]


def test_replacing_a_signature_leaves_only_one_on_disk(config):
    signature.save(png(), config)
    (config / "signature.jpg").write_bytes(b"a stale signature from before")
    signature.save(png(), config)
    assert sorted(p.name for p in config.glob("signature.*")) == ["signature.png"]


def test_removing_says_whether_there_was_one(config):
    assert signature.remove(config) is False
    signature.save(png(), config)
    assert signature.remove(config) is True
    assert signature.stored_path(config) is None


# --------------------------------------------------------------------------
# how big it prints
# --------------------------------------------------------------------------


def test_a_signature_is_scaled_to_fit_its_line():
    width, height = signature.size_for(png(900, 260))
    assert height.mm == pytest.approx(signature.MAX_HEIGHT_MM, abs=0.1)
    assert width.mm <= signature.MAX_WIDTH_MM


def test_a_long_thin_signature_is_held_back_by_its_width():
    """Scaling on height alone would print this one four inches wide and run it
    off the end of the line it is meant to sit on."""
    width, height = signature.size_for(png(2000, 120))
    assert width.mm == pytest.approx(signature.MAX_WIDTH_MM, abs=0.1)
    assert height.mm < signature.MAX_HEIGHT_MM


# --------------------------------------------------------------------------
# where it lands in a filing
# --------------------------------------------------------------------------


def paragraphs(path: Path) -> list[tuple[str, bool]]:
    """Each paragraph's text, and whether a picture was drawn in it."""
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    out = []
    for block in PARAGRAPH.findall(xml):
        text = "".join(
            "\t" if m.group(0).startswith("<w:tab") else m.group(1)
            for m in re.finditer(r'<w:tab\s*/>|<w:t(?:\s[^>]*)?>(.*?)</w:t>', block, re.S)
        )
        out.append((text, "<w:drawing>" in block))
    return out


def generate(registry, tmp_path, today, **kwargs):
    return engine.generate(
        registry, "dhs", "dhs_1p_1c", fixtures.BASE,
        today=today, output_root=tmp_path, **kwargs,
    )


def test_every_signature_sits_directly_above_the_attorneys_own_block(
    signed_registry, tmp_path, today
):
    """The check that matters. A signature over the petitioners' line, the
    notary's or the judge's would be somebody else's name signed by this
    program, so each one is required to be immediately above the block naming
    the attorney."""
    result = generate(signed_registry, tmp_path / "out", today)
    signed_anywhere = 0

    for generated in result.files:
        lines = paragraphs(generated.path)
        for i, (_text, has_image) in enumerate(lines):
            if not has_image:
                continue
            signed_anywhere += 1
            following = next(
                (t.strip() for t, _ in lines[i + 1:] if t.strip()), ""
            )
            assert "OBA #" in following, (
                f"{Path(generated.path).name}: a signature was placed above "
                f"{following!r}, which is not the attorney's block"
            )

    assert signed_anywhere >= 4, "the filing came out with almost nothing signed"


def test_nothing_is_signed_when_the_office_has_no_signature(registry, tmp_path, today):
    result = generate(registry, tmp_path / "out", today)
    for generated in result.files:
        assert not any(drawn for _text, drawn in paragraphs(generated.path))


def test_the_printed_line_comes_back_when_there_is_no_signature(registry, tmp_path, today):
    notice = next(f for f in generate(registry, tmp_path / "out", today).files
                  if f.template.endswith("notice_to_tribe.docx"))
    assert "By_________________________" in engine.document_text(notice.path)


def test_signing_can_be_turned_off_for_one_filing(signed_registry, tmp_path, today):
    """Uploading a signature is not a decision to sign everything. A filing the
    attorney means to sign by hand is an ordinary thing to want."""
    result = generate(signed_registry, tmp_path / "out", today, sign=False)
    for generated in result.files:
        assert not any(drawn for _text, drawn in paragraphs(generated.path))
    notice = next(f for f in result.files if f.template.endswith("notice_to_tribe.docx"))
    assert "By_________________________" in engine.document_text(notice.path)


def test_a_draft_is_never_signed(signed_registry, tmp_path, today):
    """A document still carrying [ Child — Date of birth ] is not one to put a
    signature on, even an unfinished one."""
    unfinished = {k: v for k, v in fixtures.BASE.items() if k != "petitioner1_address"}
    result = engine.generate(
        signed_registry, "dhs", "dhs_1p_1c", unfinished,
        today=today, output_root=tmp_path / "out", draft=True,
    )
    assert result.gaps
    for generated in result.files:
        assert not any(drawn for _text, drawn in paragraphs(generated.path))


def test_a_signature_deleted_after_startup_does_not_stop_a_filing(
    signed_registry, tmp_path, today
):
    """The path is read once at startup. If the file goes away in between, the
    filing should come out with blank signature lines rather than fail."""
    Path(signed_registry.settings[SIGNATURE_SETTING]).unlink()
    result = generate(signed_registry, tmp_path / "out", today)
    assert result.files
    for generated in result.files:
        assert not any(drawn for _text, drawn in paragraphs(generated.path))


# --------------------------------------------------------------------------
# the office details screen
# --------------------------------------------------------------------------


def test_the_signature_is_not_offered_as_a_text_box(registry, config):
    """It is a file with its own upload control. Listing it among the office's
    text fields would invite someone to type a path into it."""
    from app.registry import office_fields

    settings = {**registry.settings, SIGNATURE_SETTING: str(config / "signature.png")}
    rows = office_fields(registry.schema, settings)
    assert not any(row["id"] == SIGNATURE_SETTING for row in rows)


def test_saving_office_details_never_writes_the_signature_into_settings(config):
    from app.registry import write_settings

    path = write_settings(
        {"attorney_short_name": "Sample Attorney", SIGNATURE_SETTING: "/tmp/somewhere.png"},
        config,
    )
    written = path.read_text(encoding="utf-8")
    assert "Sample Attorney" in written
    assert SIGNATURE_SETTING not in written


def test_describe_reports_what_the_screen_needs(config):
    assert signature.describe(config) == {"present": False}
    signature.save(png(), config)
    described = signature.describe(config)
    assert described["present"] is True
    assert described["pixels"] == "900 × 260"
    assert described["name"] == "signature.png"


# --------------------------------------------------------------------------
# on the line, not in it
# --------------------------------------------------------------------------
#
# docxtpl only knows how to place a picture inline, which makes the signature a
# character on the line: the line grows to the height of the image, the block
# opens up, and the printed rule ends up under the middle of the name. Anchoring
# it with no wrapping takes it out of the text flow, so the block keeps the
# height it has unsigned and the signature lies over the rule.

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


def test_the_signature_is_anchored_rather_than_inline(signed_registry, tmp_path, today):
    result = generate(signed_registry, tmp_path / "out", today)
    drawings = 0
    for generated in result.files:
        document = Document(str(generated.path))
        for paragraph in document.paragraphs:
            for _ in paragraph._p.findall(".//" + qn("w:drawing")):
                drawings += 1
            assert not paragraph._p.findall(".//" + qn("wp:inline")), \
                "an inline picture would grow the line it sits on"
    assert drawings >= 4


def test_the_printed_rule_stays_under_the_signature(signed_registry, tmp_path, today):
    """The line prints whether or not anybody has signed, so a signed filing and
    an unsigned one are the same document with ink added."""
    signed = generate(signed_registry, tmp_path / "signed", today)
    notice = next(f for f in signed.files if f.template.endswith("notice_to_tribe.docx"))
    assert "By_________________________" in engine.document_text(notice.path)


def test_the_signature_hangs_above_the_line_it_sits_on(signed_registry, tmp_path, today):
    """Lifted by its own height less a baseline's worth, so its foot lands on the
    rule rather than the whole picture floating above or below it."""
    result = generate(signed_registry, tmp_path / "out", today)
    document = Document(str(next(f.path for f in result.files
                                 if f.template.endswith("notice_to_tribe.docx"))))
    anchors = document.element.body.findall(".//" + qn("wp:anchor"))
    assert anchors

    for anchor in anchors:
        assert anchor.findall(qn("wp:wrapNone")), "wrapping would push the text aside"
        vertical = anchor.find(qn("wp:positionV"))
        assert vertical.get("relativeFrom") == "line"
        lift = int(vertical.find(qn("wp:posOffset")).text)
        height = int(anchor.find(qn("wp:extent")).get("cy"))
        assert lift < 0, "a signature that is not lifted sits below its own line"
        assert lift == signature.BASELINE_EMU - height


def test_every_place_the_attorney_signs_has_a_line_to_sign_on(registry):
    """"By:" with nothing after it is not somewhere a pen can go — which is what
    the petitions offered, while the packet had a proper rule and two other
    templates had six underlined tabs standing in for one. Tabs only draw a line
    if the tab stops happen to be somewhere sensible, and in the pilot petition
    they were not: the stop at 720 was cleared and the only one left was at -720."""
    paragraph = re.compile(r'<w:p\b(?:(?!</w:p>).)*?</w:p>', re.S)
    offenders = []

    for matter in registry.matters:
        for variant in matter.variants:
            for document in variant.all_documents():
                path = registry.template_path(document.template)
                xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
                for block in paragraph.findall(xml):
                    text = "".join(re.findall(r'<w:t[^>]*>(.*?)</w:t>', block, re.S))
                    if "attorney_signature" not in text:
                        continue
                    if "_" * 10 not in text:
                        offenders.append(f"{path.name}: {text.strip()[:60]!r}")

    assert not offenders, "signature spots with no line:\n" + "\n".join(sorted(set(offenders)))
