"""The Filing Packet, as the office marked it up.

Eight documents in one file, so most of what matters here is consistency across
them: the same heading treatment, the same signature indent, the same checklist.
Each of these was a line on a list somebody wrote after reading a printed copy,
which is the only way some of them could have been found.
"""

from __future__ import annotations

import re
import zipfile

import pytest
from docx import Document
from docx.shared import Inches

from app import engine
from tests import fixtures

PACKET = "dhs_packet.docx"


@pytest.fixture
def packet(registry, tmp_path, today):
    """The packet as it renders with no signature on file — so the blank
    signature lines are visible rather than replaced by an image."""
    result = engine.generate(registry, "dhs", "dhs_1p_1c", fixtures.BASE,
                             today=today, output_root=tmp_path)
    generated = next(f for f in result.files if f.template.endswith(PACKET))
    return Document(str(generated.path))


def lines(document):
    return [re.sub(r"\s+", " ", p.text).strip() for p in document.paragraphs]


# --------------------------------------------------------------------------
# headings
# --------------------------------------------------------------------------

#: Each of these headings runs onto a second line. The office underlines the
#: second and not the first, which reads as one underlined phrase rather than
#: two stacked ones.
TWO_LINE_HEADINGS = [
    ("APPLICATION TO WAIVE INTERLOCUTORY DECREE", "AND WAITING PERIOD"),
    ("APPLICATION FOR BEST INTEREST AND", "FINAL DECREE OF ADOPTION"),
    ("VERIFICATION OF FILING ORIGINAL", "BIRTH CERTIFICATE OR RECORD OF BIRTH"),
]


@pytest.mark.parametrize("first,second", TWO_LINE_HEADINGS)
def test_only_the_second_line_of_a_heading_is_underlined(packet, first, second):
    underlined = {}
    for paragraph in packet.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if text in (first, second):
            underlined[text] = any(r.underline for r in paragraph.runs if r.text.strip())
    assert underlined.get(first) is False, f"{first!r} should not be underlined"
    assert underlined.get(second) is True, f"{second!r} should stay underlined"


# --------------------------------------------------------------------------
# signature blocks
# --------------------------------------------------------------------------


def test_both_submission_blocks_have_a_line_to_sign_on(packet):
    """"By:" with nothing after it is not somewhere a pen can go."""
    by_lines = [t for t in lines(packet) if t.startswith("By:")]
    assert len(by_lines) == 2
    for text in by_lines:
        assert text.count("_") >= 18, text


def test_every_block_somebody_signs_sits_at_three_inches(packet):
    """The attorney's two submission blocks, the verification, and all three
    judge's blocks. The blocks that only say who filed a page stay at the
    margin — "Filed By:" on the checklist, and the details printed under a
    judge's signature."""
    at_margin, indented = [], []
    for paragraph in packet.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if not text:
            continue
        indent = paragraph.paragraph_format.left_indent
        (indented if indent == Inches(3) else at_margin).append(text)

    for signed in ("By:", "JUDGE OF THE DISTRICT COURT"):
        assert not [t for t in at_margin if t.startswith(signed)], \
            f"{signed} left at the margin"
    assert len([t for t in indented if t == "JUDGE OF THE DISTRICT COURT"]) == 3
    # and the identification blocks are still where they were
    assert "Filed By:" in at_margin


# --------------------------------------------------------------------------
# the rest of the list
# --------------------------------------------------------------------------


def test_the_filings_checklist_is_double_spaced(packet):
    """Room to tick a box and write beside it.

    The FILINGS list only. The consent page further on has a checklist of its
    own — who is giving the consent — which was not part of the request and is
    still single spaced."""
    paragraphs = packet.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "FILINGS")
    boxes = []
    for paragraph in paragraphs[start:]:
        text = paragraph.text.strip()
        if text.startswith("☐"):
            boxes.append(paragraph)
        elif boxes and text:
            break                      # the list has ended
    assert len(boxes) == 10, [b.text for b in boxes]
    for box in boxes:
        assert box.paragraph_format.line_spacing == 2.0, box.text


def test_the_consent_checklist_was_left_alone(packet):
    """It was not on the office's list, and nothing should have changed it."""
    consent = next(p for p in packet.paragraphs
                   if p.text.strip().startswith("☐  The mother of the above-named minor"))
    assert consent.paragraph_format.line_spacing is None


def test_the_note_to_whoever_assembles_the_packet_is_not_printed(packet):
    """"[Certified copies to be attached]" is an instruction to the office, not
    something a court should be handed."""
    assert not any("Certified copies to be attached" in t for t in lines(packet))


def test_the_last_heading_leaves_room_for_the_judge_below_it(registry):
    """One blank line under it rather than two, which is what keeps the judge's
    signature block from falling onto a page of its own."""
    xml = zipfile.ZipFile(registry.template_path(f"dhs/{PACKET}")).read(
        "word/document.xml").decode("utf-8")
    paragraphs = re.findall(r'<w:p\b(?:(?!</w:p>).)*?</w:p>', xml, re.S)
    texts = ["".join(re.findall(r'<w:t[^>]*>(.*?)</w:t>', p, re.S)).strip() for p in paragraphs]

    last = max(i for i, t in enumerate(texts) if t == "STATE OF OKLAHOMA")
    blanks = 0
    while not texts[last + 1 + blanks]:
        blanks += 1
    assert blanks == 1, f"{blanks} blank lines under the last heading"
