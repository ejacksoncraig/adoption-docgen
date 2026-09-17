"""Rendering the pilot template, both sides of every conditional in it."""

from __future__ import annotations

import re
from pathlib import Path

import docx
import pytest
from jinja2 import UndefinedError

from app import engine
from app.engine import RenderError
from app.schema import IntakeError
from conftest import PILOT_MATTER, PILOT_VARIANT
from tests import fixtures


def render(registry, values, tmp_path, today, template="dhs_petition_decree_1p_1c.docx",
           matter=PILOT_MATTER, variant=PILOT_VARIANT, **kwargs) -> str:
    """Generate a variant and return the text of one of its documents.

    A variant produces several documents now — the petition and decree, the
    consent and affidavit, the filing packet — so a test has to say which one it
    is reading.
    """
    result = engine.generate(
        registry, matter, variant, values, today=today, output_root=tmp_path, **kwargs
    )
    wanted = next((f for f in result.files if f.template.endswith(template)), None)
    assert wanted is not None, f"{template} not among {[f.template for f in result.files]}"
    return engine.document_text(wanted.path)


@pytest.fixture
def base_text(registry, tmp_path, today) -> str:
    return render(registry, fixtures.BASE, tmp_path, today)


@pytest.fixture
def alternate_text(registry, tmp_path, today) -> str:
    return render(registry, fixtures.ALTERNATE, tmp_path, today)


# --------------------------------------------------------------------------
# the failure that matters most
# --------------------------------------------------------------------------


@pytest.mark.parametrize("token", ["XXX", "{{", "{%"])
def test_no_placeholder_survives_into_a_document(base_text, alternate_text, token):
    assert token not in base_text
    assert token not in alternate_text


def test_a_template_that_would_leave_a_placeholder_is_discarded(registry, tmp_path, today):
    """Rather than saving a document with XXX in it, the render fails and the
    file is removed. A discarded document cannot be filed by mistake."""
    template = tmp_path / "leftover.docx"
    document = docx.Document()
    document.add_paragraph("Comes now XXX, and respectfully alleges.")
    document.save(str(template))

    out = tmp_path / "out.docx"
    with pytest.raises(RenderError) as exc:
        engine.render_document(template, {}, out)
    assert "'XXX'" in str(exc.value)
    assert not out.exists()


# --------------------------------------------------------------------------
# both sides of every branch
# --------------------------------------------------------------------------


def test_icwa_applies(base_text):
    assert "is a member of the Sample Nation" in base_text
    assert "does apply to this proceeding" in base_text
    assert "not of Indian blood" not in base_text


def test_icwa_does_not_apply(alternate_text):
    assert "not of Indian blood" in alternate_text
    assert "do not apply to this action" in alternate_text
    assert "Sample Nation" not in alternate_text


def test_name_change_requested(base_text):
    assert "desires to change the name of SAMPLE CHILD to the name of SAMPLE NEW CHILD NAME" in base_text
    assert "shall be changed from SAMPLE CHILD to SAMPLE NEW CHILD NAME" in base_text


def test_name_change_not_requested(alternate_text):
    assert "does not desire a name change" in alternate_text
    assert "SAMPLE NEW CHILD NAME" not in alternate_text
    assert "shall be changed from" not in alternate_text


def test_parental_rights_relinquished(base_text):
    assert base_text.count("permanent relinquishment") >= 2
    assert "parental rights were terminated" not in base_text


def test_parental_rights_terminated(alternate_text):
    assert alternate_text.count("parental rights were terminated") >= 2
    assert "permanent relinquishment" not in alternate_text


# The child's consent sentence has to be matched precisely: the two biological
# parent paragraphs contain "consent to this Adoption is not required" as well.
CHILD_CONSENT = "child’s consent to this Adoption is "


def test_child_under_twelve_does_not_consent(base_text):
    assert "said child is less than 12 years of age" in base_text
    assert CHILD_CONSENT + "not required" in base_text
    assert "minor’s consent to this adoption is not required" in base_text  # the decree


def test_child_over_twelve_must_consent(alternate_text):
    assert "said child is more than 12 years of age" in alternate_text
    assert CHILD_CONSENT + "required" in alternate_text
    assert CHILD_CONSENT + "not required" not in alternate_text
    assert "minor’s consent to this adoption is required" in alternate_text


def test_mother_and_father_branch_independently(registry, tmp_path, today):
    """The two paragraphs read the same way but from different answers; a template
    edit that wired both to the mother's status would pass a single-branch test."""
    text = render(registry, fixtures.values(bio_mother_status="relinquished",
                                            bio_father_status="terminated"), tmp_path, today)
    mother = text.split("12.")[0]
    father = text.split("12.")[1]
    assert "permanent relinquishment of their parental rights" in mother
    assert "parental rights were terminated" in father


# --------------------------------------------------------------------------
# values that arrive by way of a derivation
# --------------------------------------------------------------------------


def test_dates_and_ages_are_rendered_in_the_shape_each_sentence_needs(base_text):
    assert "DOB:" in base_text and "3/2/2015" in base_text          # caption, M/D/YYYY
    assert "born on the 2nd day of March, 2015" in base_text        # prose, ordinal
    assert "is now 41 years of age" in base_text                    # computed from DOB
    assert "in September of 2024" in base_text                      # foster placement: month, not day


def test_every_county_in_the_document_follows_the_filing_county(registry, tmp_path, today):
    """The verification block used to read COUNTY OF WAGONER whatever was filed."""
    text = render(registry, fixtures.values(county="Tulsa"), tmp_path, today)
    assert "IN AND FOR TULSA COUNTY" in text
    assert "resident of Tulsa County" in text
    assert "COUNTY OF TULSA" in text
    assert "COUNTY OF WAGONER" not in text
    # "Wagoner, OK 74477" survives: that is the office's own address in the
    # signature block, which is still literal text. See the TODO in
    # docs/TEMPLATE_AUTHORING.md.


def test_the_notary_year_is_not_frozen(base_text):
    assert "day of __________________, 2026." in base_text  # TODAY is in 2026


def test_blank_case_number_prints_the_year_and_a_blank(base_text):
    assert "No. FA-2026-_____" in base_text


# --------------------------------------------------------------------------
# one document, several variants
# --------------------------------------------------------------------------
#
# The DHS consent is filed in every DHS adoption, but it names a second
# petitioner only when there is one. matters.json gives that document
# "field_groups": ["petitioner2"], and the template guards the value with
# {% if petitioner2_name %}. See GuardedUndefined in app/registry.py.

CONSENT = "dhs_consent_affidavit_order.docx"


def test_a_shared_document_omits_a_petitioner_the_variant_does_not_collect(registry, tmp_path, today):
    text = render(registry, fixtures.BASE, tmp_path, today, template=CONSENT)
    assert "adoption of the above named child by SAMPLE PETITIONER." in text
    assert "and SAMPLE" not in text.split("above named child by")[1][:80]


def test_the_same_document_names_both_petitioners_when_there_are_two(registry, tmp_path, today):
    values = {
        **fixtures.BASE,
        "petitioner2_name": "SAMPLE SECOND PETITIONER",
        "petitioner2_race": "Caucasian",
        "petitioner2_gender": "male",
        "petitioner2_dob": "1983-02-11",
        "petitioner2_birth_state": "Oklahoma",
        "marriage_date": "2014-10-11",
        "marriage_place": "Sampletown, Oklahoma",
    }
    text = render(registry, values, tmp_path, today, template=CONSENT, variant="dhs_2p_1c")
    assert "by SAMPLE PETITIONER and SAMPLE SECOND PETITIONER." in text


def test_an_optional_value_may_be_tested_but_not_printed(registry, tmp_path, today):
    """GuardedUndefined draws the line: {% if x %} is allowed, {{ x }} is not."""
    from app.registry import jinja_env

    env = jinja_env()
    assert env.from_string("{% if absent %}yes{% else %}no{% endif %}").render() == "no"
    with pytest.raises(UndefinedError):
        env.from_string("Comes now {{ absent }},").render()


# --------------------------------------------------------------------------
# refusing to generate
# --------------------------------------------------------------------------


def test_incomplete_intake_generates_nothing(registry, tmp_path, today):
    with pytest.raises(IntakeError):
        engine.generate(registry, PILOT_MATTER, PILOT_VARIANT,
                        fixtures.values(child1_name=None), today=today, output_root=tmp_path)
    assert list(tmp_path.rglob("*.docx")) == []


def test_unset_office_settings_stop_generation(registry, tmp_path, today):
    """The attorney block is not on the intake form; if settings.json is empty the
    decree would name no attorney of record."""
    saved = registry.settings.pop("attorney_short_name")
    try:
        with pytest.raises(RenderError) as exc:
            engine.generate(registry, PILOT_MATTER, PILOT_VARIANT, fixtures.BASE,
                            today=today, output_root=tmp_path)
        assert any("settings.json" in p for p in exc.value.problems)
        assert list(tmp_path.rglob("*.docx")) == []
    finally:
        registry.settings["attorney_short_name"] = saved


# --------------------------------------------------------------------------
# where the files land
# --------------------------------------------------------------------------


def test_the_saved_package_is_valid_docx(registry, tmp_path, today):
    """docxtpl leaves a repeated part in the zip; a filing has to be a clean package."""
    import collections
    import zipfile

    result = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT, fixtures.BASE,
                             today=today, output_root=tmp_path)
    names = zipfile.ZipFile(result.files[0].path).namelist()
    assert [n for n, count in collections.Counter(names).items() if count > 1] == []
    assert "word/document.xml" in names
    docx.Document(str(result.files[0].path))  # opens


def test_output_folder_and_file_names(registry, tmp_path, today):
    result = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT, fixtures.BASE,
                             today=today, output_root=tmp_path)
    assert result.folder.name == "dhs_SAMPLE_CHILD_2026-08-21"
    # Numbered so the folder sorts into filing order; the child named so a file
    # emailed on its own still says who it is about.
    assert result.files[0].path.name == "1 Petition and Decree - SAMPLE CHILD.docx"


def test_generating_twice_never_overwrites(registry, tmp_path, today):
    first = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT, fixtures.BASE,
                            today=today, output_root=tmp_path)
    second = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT, fixtures.BASE,
                             today=today, output_root=tmp_path)
    assert first.folder == second.folder
    assert first.files[0].path != second.files[0].path
    assert second.files[0].path.name.endswith("(2).docx")


def test_names_with_characters_windows_forbids(registry, tmp_path, today):
    result = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT,
                             fixtures.values(child1_name='SAMPLE "CHILD" / JR'),
                             today=today, output_root=tmp_path)
    name = result.files[0].path.name
    assert not set(name) & set('<>:"/\\|?*')
    assert result.files[0].path.exists()


def test_pdf_is_skipped_with_a_message_when_libreoffice_is_absent(registry, tmp_path, today, monkeypatch):
    monkeypatch.setattr(engine, "find_soffice", lambda: None)
    result = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT, fixtures.BASE,
                             today=today, output_root=tmp_path, pdf=True)
    assert result.files[0].pdf is None
    assert any("LibreOffice" in w for w in result.warnings)
    assert result.files[0].path.exists()  # the .docx is still there


@pytest.mark.skipif(engine.find_soffice() is None, reason="LibreOffice is not installed")
def test_pdf_export_when_libreoffice_is_present(registry, tmp_path, today):
    result = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT, fixtures.BASE,
                             today=today, output_root=tmp_path, pdf=True)
    assert result.files[0].pdf is not None
    assert result.files[0].pdf.exists()
    assert result.files[0].pdf.stat().st_size > 1000


# --------------------------------------------------------------------------
# PDF export under concurrency
# --------------------------------------------------------------------------


def test_each_pdf_conversion_gets_its_own_libreoffice_profile(monkeypatch, tmp_path):
    """LibreOffice shares one profile per account unless told otherwise, and two
    conversions running at once against it collide — the second either fails or
    silently writes nothing. Certain to happen the moment more than one person is
    served at once."""
    seen = []

    class FakeProc:
        returncode = 0
        stdout = stderr = ""

    def fake_run(command, **_kwargs):
        seen.append(command)
        Path(command[-1]).with_suffix(".pdf").write_bytes(b"%PDF-1.4\n")
        return FakeProc()

    monkeypatch.setattr(engine.subprocess, "run", fake_run)

    for name in ("one.docx", "two.docx"):
        document = tmp_path / name
        docx.Document().save(str(document))
        engine.export_pdf(document, soffice="soffice")

    profiles = [
        argument for command in seen for argument in command
        if argument.startswith("-env:UserInstallation=")
    ]
    assert len(profiles) == 2, "every conversion must name a profile"
    assert profiles[0] != profiles[1], "two conversions shared a profile"
    assert all(p.startswith("-env:UserInstallation=file://") for p in profiles)


def test_a_conversion_profile_does_not_outlive_the_conversion(monkeypatch, tmp_path):
    """A throwaway profile is a few MB; leaving one behind per PDF would fill a
    server's disk quietly."""
    captured = {}

    class FakeProc:
        returncode = 0
        stdout = stderr = ""

    def fake_run(command, **_kwargs):
        captured["profile"] = next(
            a.split("=", 1)[1] for a in command if a.startswith("-env:UserInstallation=")
        )
        Path(command[-1]).with_suffix(".pdf").write_bytes(b"%PDF-1.4\n")
        return FakeProc()

    monkeypatch.setattr(engine.subprocess, "run", fake_run)
    document = tmp_path / "one.docx"
    docx.Document().save(str(document))
    engine.export_pdf(document, soffice="soffice")

    from urllib.parse import urlparse
    from urllib.request import url2pathname

    left = Path(url2pathname(urlparse(captured["profile"]).path))
    assert not left.exists()


# --------------------------------------------------------------------------
# the petitioners' names are not set apart
# --------------------------------------------------------------------------


def _bold_name_paragraphs(path):
    """Paragraphs where a petitioner's name is bold, and whether the whole
    paragraph is."""
    import zipfile

    paragraph = re.compile(r'<w:p\b(?:(?!</w:p>).)*?</w:p>', re.S)
    run = re.compile(r'<w:r>(?:(?!</w:r>).)*?</w:r>', re.S)
    name = re.compile(r'\{\{\s*petitioner[12]_name\s*\}\}')

    def text_of(one):
        return "".join(
            "\t" if m.group(0).startswith("<w:tab") else m.group(1)
            for m in re.finditer(r'<w:tab\s*/>|<w:t(?:\s[^>]*)?>(.*?)</w:t>', one, re.S))

    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    for block in paragraph.findall(xml):
        speaking = [r for r in run.findall(block) if text_of(r).strip()]
        if not any(name.search(text_of(r)) and "<w:b/>" in r for r in speaking):
            continue
        yield all("<w:b/>" in r for r in speaking), text_of(speaking[0])


PETITIONS = [
    "dhs/dhs_petition_decree_1p_1c.docx",
    "dhs/dhs_petition_decree_2p_1c.docx",
    "dhs/dhs_petition_1p_2c.docx",
    "dhs/dhs_petition_2p_2c.docx",
    "stepparent/step_petition_1c.docx",
    "stepparent/step_petition_2c.docx",
]


@pytest.mark.parametrize("template", PETITIONS)
def test_a_petitioners_name_is_not_bold_in_the_middle_of_a_sentence(registry, template):
    """It used to be, in every petition, which set the name apart from the
    sentence it belongs to for no reason anybody could give."""
    offenders = [opening for whole, opening in _bold_name_paragraphs(
        registry.template_path(template)) if not whole]
    assert not offenders, "\n".join(offenders)


def test_a_wholly_bold_paragraph_keeps_its_bold(registry):
    """The decree's two ordering paragraphs are bold from end to end. There the
    bold belongs to the paragraph, not the name, and lifting it off the name
    alone would punch a hole in it."""
    whole = [opening for is_whole, opening in _bold_name_paragraphs(
        registry.template_path("dhs/dhs_petition_decree_1p_1c.docx")) if is_whole]
    assert len(whole) == 2
    assert all(t.startswith("IT IS ") for t in whole), whole


# --------------------------------------------------------------------------
# the Verification stands on a page of its own
# --------------------------------------------------------------------------


VERIFICATION_TEMPLATES = [
    "dhs/dhs_petition_decree_1p_1c.docx",
    "dhs/dhs_petition_decree_2p_1c.docx",
    "dhs/dhs_petition_1p_2c.docx",
    "dhs/dhs_petition_2p_2c.docx",
    "stepparent/step_petition_1c.docx",
    "stepparent/step_petition_2c.docx",
]


def _verification_paragraph(document):
    from docx.oxml.ns import qn

    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "VERIFICATION":
            return paragraph, paragraph._p.find(qn("w:pPr"))
    raise AssertionError("no VERIFICATION heading in the document")


@pytest.mark.parametrize("template", VERIFICATION_TEMPLATES)
def test_the_verification_starts_a_new_page(registry, template):
    """It used to be reached by blank lines, which is not an instruction — it is
    room left in the hope that the text above happens to end in the right place.
    How long a petition runs depends on ICWA, a name change, kinship and how many
    children there are, so the hope was misplaced."""
    import docx
    from docx.oxml.ns import qn

    document = docx.Document(str(registry.template_path(template)))
    _paragraph, properties = _verification_paragraph(document)
    assert properties is not None, f"{template}: VERIFICATION has no paragraph properties"
    assert properties.find(qn("w:pageBreakBefore")) is not None, template


@pytest.mark.parametrize("branches", [
    {"icwa_applies": True, "name_change": True, "kinship": True,
     "kinship_relationship": "maternal grandparents", "child1_new_name": "NEW NAME"},
    {"icwa_applies": False, "name_change": False, "tribe": None,
     "child1_new_name": None, "kinship": False},
    {"bio_mother_status": "terminated", "bio_father_status": "terminated"},
    {"child1_dob": "2008-01-01"},
])
def test_the_verification_keeps_its_page_whatever_the_petition_says(
    registry, tmp_path, today, branches
):
    """The point of a page break rather than padding: the length of everything
    above it stops mattering."""
    import docx
    from docx.oxml.ns import qn

    result = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT,
                             fixtures.values(**branches), today=today, output_root=tmp_path)
    generated = next(f for f in result.files if "petition_decree" in f.template)
    _paragraph, properties = _verification_paragraph(docx.Document(str(generated.path)))
    assert properties.find(qn("w:pageBreakBefore")) is not None


def test_the_verification_is_not_split_across_two_pages(registry):
    """A page break says where it starts. Keeping its lines together is what
    stops Word ending the page halfway down it."""
    import docx
    from docx.oxml.ns import qn

    document = docx.Document(
        str(registry.template_path("dhs/dhs_petition_decree_1p_1c.docx")))
    started = False
    held = 0
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "VERIFICATION":
            started = True
        if not started:
            continue
        properties = paragraph._p.find(qn("w:pPr"))
        if properties is not None and properties.find(qn("w:keepLines")) is not None:
            held += 1
        if paragraph.text.strip() == "Notary Public":
            break
    assert held >= 15, f"only {held} paragraphs of the verification are held together"
