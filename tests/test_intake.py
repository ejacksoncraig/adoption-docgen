"""Saved intake files and the questionnaire emailed to adoptive parents."""

from __future__ import annotations

import json

import pytest

from app import engine, intake
from conftest import PILOT_MATTER, PILOT_VARIANT
from tests import fixtures


# --------------------------------------------------------------------------
# save / load
# --------------------------------------------------------------------------


def test_save_and_reopen_round_trips(registry, tmp_path):
    path = intake.save_intake(PILOT_MATTER, PILOT_VARIANT, fixtures.BASE, directory=tmp_path)
    loaded = intake.load_intake(path, registry)

    assert loaded["matter"] == PILOT_MATTER
    assert loaded["variant"] == PILOT_VARIANT
    assert loaded["values"] == fixtures.BASE
    assert loaded["warnings"] == []
    assert loaded["saved_at"]


def test_saved_file_is_named_for_the_child_and_the_date(tmp_path):
    path = intake.save_intake(PILOT_MATTER, PILOT_VARIANT, fixtures.BASE, directory=tmp_path)
    assert path.parent == tmp_path
    assert path.name.startswith("dhs_SAMPLE_CHILD_")
    assert path.suffix == ".json"


def test_a_field_that_no_longer_exists_is_dropped_and_reported(registry, tmp_path):
    """Loading an intake saved before a config change must not look complete when
    it is not."""
    path = tmp_path / "old.json"
    path.write_text(json.dumps({
        "format": 1, "matter": PILOT_MATTER, "variant": PILOT_VARIANT,
        "values": {**fixtures.BASE, "petitioner1_favourite_colour": "green"},
    }), encoding="utf-8")

    loaded = intake.load_intake(path, registry)
    assert "petitioner1_favourite_colour" not in loaded["values"]
    assert any("petitioner1_favourite_colour" in w for w in loaded["warnings"])


def test_an_intake_for_a_variant_that_no_longer_exists_is_refused(registry, tmp_path):
    path = tmp_path / "gone.json"
    path.write_text(json.dumps({"matter": "dhs", "variant": "dhs_9p_9c", "values": {}}), encoding="utf-8")
    with pytest.raises(ValueError, match="no longer offers"):
        intake.load_intake(path, registry)


def test_a_file_that_is_not_an_intake_is_refused(registry, tmp_path):
    path = tmp_path / "notes.json"
    path.write_text('{"hello": "world"}', encoding="utf-8")
    with pytest.raises(ValueError, match="not an intake file"):
        intake.load_intake(path, registry)

    broken = tmp_path / "broken.json"
    broken.write_text("{ not json", encoding="utf-8")
    with pytest.raises(ValueError, match="not a valid intake file"):
        intake.load_intake(broken, registry)


def test_a_reopened_intake_still_generates(registry, tmp_path, today):
    path = intake.save_intake(PILOT_MATTER, PILOT_VARIANT, fixtures.BASE, directory=tmp_path)
    loaded = intake.load_intake(path, registry)
    result = engine.generate(registry, PILOT_MATTER, PILOT_VARIANT, loaded["values"],
                             today=today, output_root=tmp_path)
    assert result.files[0].path.exists()


def test_listing_shows_the_newest_first(tmp_path):
    import os
    import time

    first = intake.save_intake(PILOT_MATTER, PILOT_VARIANT, fixtures.values(child1_name="SAMPLE ONE"),
                               directory=tmp_path)
    second = intake.save_intake(PILOT_MATTER, PILOT_VARIANT, fixtures.values(child1_name="SAMPLE TWO"),
                                directory=tmp_path)
    os.utime(second, (time.time() + 10, time.time() + 10))

    listed = intake.list_intakes(tmp_path)
    assert [entry["name"] for entry in listed] == [second.name, first.name]
    assert listed[0]["variant"] == PILOT_VARIANT


# --------------------------------------------------------------------------
# the questionnaire
# --------------------------------------------------------------------------


def form_text(path) -> str:
    """Everything a person would read, tables included.

    engine.document_text walks body paragraphs, and the one-page form puts
    almost all of itself inside a table — so a test that used it would pass by
    finding nothing."""
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            lines.append(Paragraph(child, document).text)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, document).rows:
                lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


@pytest.fixture
def questionnaire_text(registry, tmp_path) -> str:
    path = intake.build_questionnaire(registry, PILOT_MATTER, PILOT_VARIANT, tmp_path / "q.docx")
    return form_text(path)


def test_questionnaire_asks_for_every_answer_the_family_can_give(registry, questionnaire_text):
    groups = registry.variant(PILOT_MATTER, PILOT_VARIANT).field_groups
    asked = [fd for fd in registry.schema.input_fields(groups) if registry.schema.on_questionnaire(fd)]
    assert asked
    for fd in asked:
        printed = fd.short_label or fd.label
        assert printed in questionnaire_text, f"{fd.id} is not on the questionnaire"


def test_questionnaire_leaves_out_what_the_office_fills_in(questionnaire_text):
    assert "Filing county" not in questionnaire_text      # whole group opted out
    assert "Case number" not in questionnaire_text
    assert "Fees and costs to date" not in questionnaire_text  # single field opted out


def test_questionnaire_never_asks_for_a_derived_value(registry, questionnaire_text):
    for fd in registry.schema.fields.values():
        if fd.derived and fd.label != fd.id:
            assert fd.label not in questionnaire_text


def test_questionnaire_offers_the_choices_for_a_select(questionnaire_text):
    assert "Relinquished  /  Terminated" in questionnaire_text
    # a yes/no question is "Y / N" on the form, which is what gets circled
    assert "Y  /  N" in questionnaire_text


def test_questionnaire_is_blank(questionnaire_text):
    """It is emailed out; it must not carry anyone's answers."""
    for value in ("SAMPLE CHILD", "SAMPLE PETITIONER", "Sample Nation", "2015"):
        assert value not in questionnaire_text


# --------------------------------------------------------------------------
# placeholder data
# --------------------------------------------------------------------------


def test_sample_values_fill_every_required_answer(registry, pilot_groups):
    values = intake.sample_values(registry.schema, pilot_groups)
    assert registry.schema.missing_required(values, pilot_groups) == []
    assert registry.schema.validate(values, pilot_groups) == []


def test_sample_values_take_the_other_branch_when_asked(registry, pilot_groups):
    other = intake.sample_values(registry.schema, pilot_groups, truthy=False)
    assert other["icwa_applies"] is False
    assert other["name_change"] is False
    assert other["bio_mother_status"] == "terminated"
    assert registry.schema.missing_required(other, pilot_groups) == []


# --------------------------------------------------------------------------
# the office's own worksheet
# --------------------------------------------------------------------------
#
# The questionnaire above is the family's and leaves out what the office fills
# in for itself. The worksheet is the other half: the whole intake, on paper,
# for an attorney to take an interview on and type up afterwards.


@pytest.fixture
def worksheet_text(registry, tmp_path) -> str:
    path = intake.build_intake_worksheet(registry, PILOT_MATTER, PILOT_VARIANT, tmp_path / "w.docx")
    return form_text(path)


def test_the_worksheet_asks_every_question_the_variant_collects(registry, worksheet_text):
    """Including the ones the family is never asked. Anything missing here is a
    question somebody has to remember on their own."""
    groups = registry.variant(PILOT_MATTER, PILOT_VARIANT).field_groups
    for fd in registry.schema.input_fields(groups):
        printed = fd.short_label or fd.label
        assert printed in worksheet_text, f"{fd.id} is not on the worksheet"


def test_the_worksheet_carries_what_the_questionnaire_leaves_out(registry, tmp_path, worksheet_text):
    questionnaire = engine.document_text(
        intake.build_questionnaire(registry, PILOT_MATTER, PILOT_VARIANT, tmp_path / "q2.docx"))
    for office_only in ("Filing county", "Case number", "Fees and costs to date"):
        assert office_only in worksheet_text, office_only
        assert office_only not in questionnaire, office_only


def test_the_worksheet_says_which_questions_are_conditional(worksheet_text):
    """A blank on paper is ambiguous — not asked, or asked and answered "no"?

    The long form said so in words under each question. One page has no room for
    that, so the question carries a dagger and the page carries one footnote."""
    assert "Tribe\u2020:" in worksheet_text
    assert "asked only when the box it follows is ticked" in worksheet_text


def test_the_worksheet_offers_somewhere_to_write_every_answer(registry, tmp_path):
    """A question with no line under it is a question that gets skipped.

    On the one-page form the space is a ruled table cell rather than a run of
    underscores, except where the answer is a choice to circle."""
    import docx
    from docx.oxml.ns import qn

    path = intake.build_intake_worksheet(registry, PILOT_MATTER, PILOT_VARIANT, tmp_path / "w3.docx")
    table = docx.Document(str(path)).tables[0]

    labelled = {}
    for row in table.rows:
        cells = list(row.cells)
        for label_cell, answer_cell in ((cells[0], cells[1]), (cells[2], cells[3])):
            text = label_cell.text.strip().rstrip(":").rstrip("\u2020")
            if not text or text.isupper():
                continue
            ruled = answer_cell._tc.findall(".//" + qn("w:bottom"))
            labelled[text] = bool(ruled) or bool(answer_cell.text.strip())

    for fd in registry.schema.input_fields(
            registry.variant(PILOT_MATTER, PILOT_VARIANT).field_groups):
        printed = (fd.short_label or fd.label)
        assert labelled.get(printed), f"{fd.id} has nowhere to write"


def test_every_ready_variant_produces_a_worksheet(registry, tmp_path):
    for matter in registry.matters:
        for variant in matter.variants:
            if not variant.is_ready:
                continue
            path = intake.build_intake_worksheet(
                registry, matter.id, variant.id, tmp_path / f"{variant.id}.docx")
            assert path.exists() and path.stat().st_size > 0
